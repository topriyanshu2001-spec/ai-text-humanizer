# -*- coding: utf-8 -*-
import copy, re, sys
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
sys.path.insert(0, ".")
from rewrite_m2 import R

SRC = "/root/.claude/uploads/82d0c4b6-3f37-59a9-b0b3-0bd39093c6a3/22ae8ba1-Migraine_Beyond_CGRP_Ardhavabhedaka_Review.docx"
EN = "–"

def uni(t):
    t = t.replace("[8-10]", "[8" + EN + "10]")
    t = t.replace("1.04-3.09", "1.04" + EN + "3.09")
    t = t.replace("between 1 and 2.5 days", "between 1 and 2.5 days")
    t = t.replace("pancreatic b-cells", "pancreatic β-cells")
    t = t.replace("NF-kB", "NF-κB")
    return re.sub(r"[ \t]{2,}", " ", t).strip()

d = docx.Document(SRC)
paras = list(d.paragraphs)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""

rep = ins = dele = 0
for idx in sorted(R, reverse=True):
    new = R[idx]
    p = paras[idx]
    assert "<w:drawing>" not in p._p.xml, f"image paragraph {idx}"
    if not new:
        p._p.getparent().remove(p._p); dele += 1; continue
    set_text(p, uni(new[0])); rep += 1
    anchor = p._p
    for extra in new[1:]:
        clone = copy.deepcopy(p._p)
        anchor.addnext(clone); anchor = clone
        set_text(Paragraph(clone, p._parent), uni(extra)); ins += 1

# Remove the two illustration-prompt boxes. Their own first line says
# "Delete this box before submission"; they are literal image-generation prompts.
removed_boxes = 0
for tbl in list(d.tables):
    txt = "\n".join(c.text for r in tbl.rows for c in r.cells)
    if "Illustration prompt for Figure" in txt and "Delete this box before submission" in txt:
        tbl._tbl.getparent().remove(tbl._tbl); removed_boxes += 1

out = "Migraine_Beyond_CGRP_Ardhavabhedaka_Review_v2.docx"
d.save(out)
print(f"replaced {rep}, inserted {ins}, deleted {dele} paragraphs; "
      f"removed {removed_boxes} illustration-prompt boxes -> {out}")

d2 = docx.Document(out)
W = lambda t: len(re.findall(r"[A-Za-z][A-Za-z'\-]*", t))
ps = [p.text for p in d2.paragraphs]
iref = next(i for i, t in enumerate(ps) if t.strip() == "References")
main = sum(W(t) for t in ps[:iref])
print(f"paragraphs {len(ps)}, tables {len(d2.tables)}, words before References: {main}")
