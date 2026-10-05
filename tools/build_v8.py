# -*- coding: utf-8 -*-
import json, copy, re, sys
import docx
sys.path.insert(0, ".")
from rewrite_v8 import R

SRC = "/root/.claude/uploads/82d0c4b6-3f37-59a9-b0b3-0bd39093c6a3/701cfe51-Rasayana_DushiVisha_v4.docx"
V7 = [p["txt"] for p in json.load(open("v7_paras.json"))]

NBH = "‑"; EN = "–"; MU = "µ"; KAPPA = "κ"

def uni(t):
    """Restore the manuscript's own Unicode conventions."""
    t = t.replace("ug/dL", MU + "g/dL")
    t = t.replace("NF-kB", "NF" + NBH + KAPPA + "B")
    t = t.replace("AP-1", "AP" + NBH + "1")
    t = t.replace("gut-brain", "gut" + EN + "brain")
    t = t.replace("blood-brain", "blood" + EN + "brain")
    t = t.replace("herb-drug", "herb" + EN + "drug")
    t = t.replace("[2-4]", "[2" + EN + "4]")
    return t

def clean(t):
    """Repair PDF-extraction artifacts in transplanted v7 text."""
    t = t.replace("timeand dose", "time- and dose")
    t = re.sub(r"\s+([,;.])", r"\1", t)            # "monnieri , one" -> "monnieri, one"
    t = re.sub(r"[ \t]{2,}", " ", t).strip()
    return t

def prep(t):
    return clean(uni(t))

# v7[106] has the next heading glued on by the PDF extractor
SAFETY = V7[106].replace("Future directions and translational strategy", "").strip()

# docx paragraph index -> list of replacement paragraphs ([] deletes, None leaves alone)
PLAN = {
    22:  R[3],                       # ABSTRACT
    39:  [R[6][0]],                  # Intro p1 (specific half)
    40:  [R[6][1]],                  # Intro p1 (general half, split off)
    41:  R[7],                       # Intro p2 -> 2 paragraphs
    42:  [V7[8]],
    47:  [V7[12]], 48: [V7[13]],
    50:  [V7[15]], 51: [V7[16]],
    53:  [V7[18]],
    54:  R[19],                      # heading rename
    55:  [R[20][0]], 56: [R[20][1]],
    59:  [V7[22]],
    61:  [V7[24]], 62: [V7[25]],
    64:  R[27], 65: [],              # Guduchi summary paragraph absorbed
    68:  R[30], 69: [V7[31]],
    71:  [V7[33]],
    73:  [V7[35]], 74: [V7[36]],
    79:  [V7[66]], 81: [V7[68]],
    84:  R[70],                      # sub-heading rename
    85:  R[71],
    87:  R[73],
    91:  [R[77][0]], 92: [R[77][1]],
    98:  [SAFETY],
    101: R[108], 102: [],            # tail absorbed
    103: R[109],                     # heading rename
    104: R[110],                     # numbered list -> prose
    105: [], 106: [], 107: [],
    108: [V7[114]], 109: [],         # merged
    111: [R[116][0]], 112: [R[116][1]],
}

d = docx.Document(SRC)
paras = list(d.paragraphs)

def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""

inserted = deleted = replaced = 0
for idx in sorted(PLAN, reverse=True):
    new = PLAN[idx]
    if new is None:
        continue
    p = paras[idx]
    assert "<w:drawing>" not in p._p.xml, f"refusing to touch image paragraph {idx}"
    if not new:
        p._p.getparent().remove(p._p); deleted += 1; continue
    set_text(p, prep(new[0])); replaced += 1
    anchor = p._p
    for extra in new[1:]:
        clone = copy.deepcopy(p._p)
        anchor.addnext(clone)
        anchor = clone
        np = docx.text.paragraph.Paragraph(clone, p._parent)
        set_text(np, prep(extra)); inserted += 1

out = "Rasayana_DushiVisha_v8.docx"
d.save(out)
print(f"replaced {replaced}, inserted {inserted}, deleted {deleted} -> {out}")

d2 = docx.Document(out)
print("paragraphs:", len(d2.paragraphs), " tables:", len(d2.tables),
      " drawings:", d2.element.body.xml.count("<w:drawing>"))
txt = "\n".join(p.text for p in d2.paragraphs)
print("words:", len(re.findall(r"[A-Za-z][A-Za-z'\-]*", txt)))
open("v8_docx_text.txt", "w").write(txt)
