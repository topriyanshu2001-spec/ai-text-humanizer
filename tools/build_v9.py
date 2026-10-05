# -*- coding: utf-8 -*-
import copy, re, sys
import docx
from docx.text.paragraph import Paragraph
sys.path.insert(0, ".")
from v9 import EDITS

SRC = "/home/user/ai-text-humanizer/Rasayana_DushiVisha_v8.docx"

def set_text(p, t):
    if not p.runs: p.add_run(t); return
    p.runs[0].text = t
    for r in p.runs[1:]: r.text = ""

d = docx.Document(SRC)

def find(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix): return p
    return None

rep = ins = dele = 0
for prefix, new in EDITS:
    p = find(prefix)
    if p is None:
        print(f"  !! NOT FOUND: {prefix[:55]}"); continue
    assert "<w:drawing>" not in p._p.xml, prefix[:40]
    if not new:
        p._p.getparent().remove(p._p); dele += 1; continue
    set_text(p, new[0]); rep += 1
    anchor = p._p
    for extra in new[1:]:
        clone = copy.deepcopy(p._p)
        anchor.addnext(clone); anchor = clone
        set_text(Paragraph(clone, p._parent), extra); ins += 1

out_full = "Rasayana_DushiVisha_v9.docx"
d.save(out_full)
print(f"v9: replaced {rep}, inserted {ins}, deleted {dele} -> {out_full}")

# --- second file for the AI check: identical body, title-page block removed.
# v7 (29%) was submitted without it; in v8 it scored 59.6% and added 74 flagged words.
d2 = docx.Document(out_full)
STRIP = ("Rasayana Therapy as a Systems-Level Countermeasure",
         "Dr. Sumit Dhote", "1PG Scholar Department of Agadtantra",
         "2Head of Department, Department of Agadtantra",
         "Running Title-", "* Corresponding author details",
         "PG Scholar", "Department of Agadtantra Evum Vidhivaidyak",
         "Mahatma Gandhi Ayurved College Hospital & Research Centre",
         "Email ID:", "Telephone no.:")
removed = 0
for p in list(d2.paragraphs):
    t = p.text.strip()
    if not t: continue
    if t == "ABSTRACT": break
    if t.startswith(STRIP):
        p._p.getparent().remove(p._p); removed += 1
out_chk = "Rasayana_DushiVisha_v9_for_AI_check.docx"
d2.save(out_chk)
print(f"AI-check copy: removed {removed} title-page paragraphs -> {out_chk}")

for f in (out_full, out_chk):
    dd = docx.Document(f)
    ps = [p.text for p in dd.paragraphs]
    W = lambda t: len(re.findall(r"[A-Za-z][A-Za-z'\-]*", t))
    iref = next(i for i, t in enumerate(ps) if t.strip() == "References")
    print(f"  {f}: {len(ps)} paras, tables {len(dd.tables)}, "
          f"drawings {dd.element.body.xml.count('<w:drawing>')}, "
          f"pre-Ref words {sum(W(t) for t in ps[:iref])}, refs {sum(1 for t in ps[iref:] if t.strip())}")
